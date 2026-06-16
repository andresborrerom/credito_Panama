"""Compilador del deck Producto B — Retornos esperados del portafolio LUZ.

Arma un .pptx 16:9 con portada + slides de narrativa + láminas full-bleed de
los PNGs ya generados. Auto-descubre los outputs disponibles, de modo que el
deck refleja el ESTADO ACTUAL del pipeline (se puede re-correr a medida que
se generan más cortes).

Estructura:
  0  Portada
  1  Metodología (FDP predictiva vs histórica · 5 zonas)
  2..  Histograma predictivo LUZ (nube MC + HDIs + VaR/CVaR)
  N..  Sección: comparación modelo vs historia — portafolio LUZ
  M..  Sección: comparación por índice principal
  K  Matriz de señales (si existe SENALES.md)
  Z  Disclaimer / compliance

Uso:
    PYTHONPATH=src python -m tasas_mercantil.producto_b.deck
"""
from __future__ import annotations
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SLIDE_W_IN = 13.33
SLIDE_H_IN = 7.50

BLUE = RGBColor(0x2A, 0x6F, 0xB3)
DARKBLUE = RGBColor(0x1A, 0x3A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE8, 0xEF, 0xF7)
GREY = RGBColor(0x66, 0x66, 0x66)

OUT_DIR = Path("docs/producto_b/outputs")
CMP_DIR = OUT_DIR / "comparacion_historica"
DISCLAIMER = ("Documento informativo con fines analíticos. No constituye "
              "recomendación de inversión.")


# ---------------------------------------------------------------------------
# Slides nativos
# ---------------------------------------------------------------------------
def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide, color):
    bg = slide.shapes.add_shape(1, Inches(0), Inches(0),
                                Inches(SLIDE_W_IN), Inches(SLIDE_H_IN))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def _text(slide, left, top, w, h, text, size, color, bold=False,
          italic=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = text
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def _cover(prs, as_of):
    s = _blank(prs)
    _bg(s, BLUE)
    _text(s, 0.8, 2.3, SLIDE_W_IN - 1.6, 1.3,
          "Producto B — Retornos esperados", 44, WHITE, bold=True)
    _text(s, 0.8, 3.4, SLIDE_W_IN - 1.6, 0.9,
          "Portafolio LUZ · sensibilidad y señal vs historia", 26, WHITE)
    _text(s, 0.8, 4.4, SLIDE_W_IN - 1.6, 0.7,
          f"Corte de análisis · {as_of.isoformat()}  ·  horizonte 12 meses",
          20, LIGHT)
    _text(s, 0.8, 6.5, SLIDE_W_IN - 1.6, 0.5,
          "Modelo Mercantil · NN + BMA · datos EODHD/FRED/Bloomberg",
          12, LIGHT, italic=True)


def _section(prs, title, subtitle=""):
    s = _blank(prs)
    _bg(s, DARKBLUE)
    _text(s, 0.8, 2.9, SLIDE_W_IN - 1.6, 1.1, title, 34, WHITE, bold=True)
    if subtitle:
        _text(s, 0.8, 4.0, SLIDE_W_IN - 1.6, 1.5, subtitle, 18, LIGHT)


def _bullets(prs, title, bullets):
    s = _blank(prs)
    _text(s, 0.7, 0.5, SLIDE_W_IN - 1.4, 0.9, title, 28, DARKBLUE, bold=True)
    box = s.shapes.add_textbox(Inches(0.9), Inches(1.6),
                               Inches(SLIDE_W_IN - 1.8), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (txt, lvl) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if lvl == 0 else "   – ") + txt
        p.level = lvl
        r = p.runs[0]
        r.font.size = Pt(18 if lvl == 0 else 15)
        r.font.bold = (lvl == 0)
        r.font.color.rgb = DARKBLUE if lvl == 0 else GREY
        p.space_after = Pt(8)
    _text(s, 0.7, 7.0, SLIDE_W_IN - 1.4, 0.4, DISCLAIMER, 9, GREY, italic=True)


def _image(prs, png, header):
    s = _blank(prs)
    _text(s, 0.4, 0.12, SLIDE_W_IN - 0.8, 0.4, header, 14, DARKBLUE, bold=True)
    s.shapes.add_picture(str(png), Inches(0.4), Inches(0.6),
                         width=Inches(SLIDE_W_IN - 0.8))


def _signals_table(prs, md_path: Path):
    """Construye una tabla nativa desde SENALES.md (si existe)."""
    if not md_path.exists():
        return False
    lines = [l for l in md_path.read_text(encoding="utf-8").splitlines()
             if l.strip().startswith("|")]
    if len(lines) < 2:
        return False
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines]
    header = rows[0]
    data = [r for r in rows[2:]]  # saltar separador
    if not data:
        return False

    s = _blank(prs)
    _text(s, 0.7, 0.4, SLIDE_W_IN - 1.4, 0.7,
          "Matriz de señales — modelo vs historia", 26, DARKBLUE, bold=True)
    nrows, ncols = len(data) + 1, len(header)
    tbl = s.shapes.add_table(nrows, ncols, Inches(0.6), Inches(1.4),
                             Inches(SLIDE_W_IN - 1.2), Inches(5.0)).table
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = h.replace("**", "")
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        c.text_frame.paragraphs[0].runs[0].font.bold = True
        c.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
        c.fill.solid(); c.fill.fore_color.rgb = BLUE
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = val.replace("**", "")
            run = c.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(11)
            up = val.upper()
            if "POSITIVA" in up:
                c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xDF, 0xF0, 0xDF)
            elif "NEGATIVA" in up:
                c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xF6, 0xDE, 0xDE)
    _text(s, 0.7, 6.9, SLIDE_W_IN - 1.4, 0.4, DISCLAIMER, 9, GREY, italic=True)
    return True


# ---------------------------------------------------------------------------
# Compilador
# ---------------------------------------------------------------------------
def _find(patterns, base):
    out = []
    for pat in patterns:
        out.extend(sorted(base.glob(pat)))
    return out


def compile_deck_producto_b(as_of: date | None = None,
                            output_path: Path | str | None = None) -> tuple:
    as_of = as_of or date.today()
    if output_path is None:
        output_path = OUT_DIR / f"deck_producto_b_{as_of.isoformat()}.pptx"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    included = []

    _cover(prs, as_of)

    _bullets(prs, "Metodología en una lámina", [
        ("Predecimos la distribución completa (FDP) de retornos del "
         "portafolio LUZ a 12 meses, no un punto.", 0),
        ("Partimos la FDP en 5 zonas con probabilidad:", 0),
        ("Riesgo (cola izq 2.5%) · Bajista · Esperado (HDI 50%) · "
         "Alcista · Sorpresa (cola der 2.5%).", 1),
        ("Contrastamos la FDP del modelo contra la FDP histórica observada "
         "(mismas ponderaciones actuales).", 0),
        ("La señal sale de la diferencia: centro por encima/debajo de la "
         "historia, más/menos confianza, más/menos riesgo de cola.", 1),
        ("Walk-forward valida qué tanto le pega el modelo a lo realizado.", 0),
    ])

    # Histograma predictivo (FDP del modelo)
    histo = _find(["luz_histograma_*.png", "luz_ejemplo*_*.png"], OUT_DIR)
    if histo:
        _section(prs, "FDP predictiva del portafolio LUZ",
                 "Nube Monte Carlo · HDIs 50/80/95 · VaR y CVaR 95")
        for p in histo:
            _image(prs, p, f"Histograma LUZ · {p.stem}")
            included.append(p.name)

    # Comparación portafolio
    port = _find(["luz_portafolio_*.png"], CMP_DIR)
    if port:
        _section(prs, "Modelo vs historia — Portafolio LUZ",
                 "FDP predictiva vs FDP histórica · señal por corte")
        for p in port:
            _image(prs, p, f"Comparación portafolio · {p.stem}")
            included.append(p.name)

    # Comparación por índice
    idx = _find(["indice_*.png"], CMP_DIR)
    if idx:
        _section(prs, "Modelo vs historia — Índices principales",
                 "LQD · IGOV · GHYG · EMB · ACWI")
        for p in idx:
            _image(prs, p, f"Comparación índice · {p.stem}")
            included.append(p.name)

    # Matriz de señales
    _signals_table(prs, CMP_DIR / "SENALES.md")

    # Cierre / compliance
    _section(prs, "Compliance", DISCLAIMER)

    prs.save(str(output_path))
    return output_path, included


if __name__ == "__main__":
    out, inc = compile_deck_producto_b()
    print(f"Deck: {out}")
    print(f"Láminas de PNG incluidas ({len(inc)}):")
    for n in inc:
        print("  -", n)
